import streamlit as st
import datetime
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Generador de Planeaciones NEM", layout="wide")

st.title("📚 Sistema de Planeaciones Escolares - Fase 5 (6° Grado)")
st.markdown("### Escuela Primaria 'Revolución Social de 1910' (CCT: 10EPR0361N)")
st.markdown("---")

# 1. Configuración de parámetros y selector de base estricta
campo_form = st.selectbox("Campo Formativo:", [
    "Lenguajes", 
    "Saberes y Pensamiento Científico", 
    "Ética, Naturaleza y Sociedades", 
    "De lo Humano y lo Comunitario"
])

# Selector de origen de la información para la planeación estricta
modo_estricto_fuente = st.radio(
    "🔒 Selecciona la base estricta para la planeación:",
    [
        "Solo Programa Analítico (Focalizado en la escuela y contexto local)",
        "Solo Programa Sintético (Contenidos y PDA nacionales de la SEP)",
        "Ambos (Integración del Programa Analítico y Sintético)"
    ]
)

proyecto_ltg = st.text_input("Proyecto vinculado de los LTG (CONALITEG):", value="La cápsula del tiempo")
periodo_val = st.text_input("Temporalidad (Ej. Semanas 8 y 9 - Octubre):", value="Semanas 8 y 9 de Octubre (19 a 23 y 26 a 30)")

# 2. Función para generar el documento de Word según la fuente elegida de forma estricta
def generar_documento_estricto(campo, proyecto, periodo, fuente):
    doc = Document()
    
    # Encabezado institucional fijo (Programa Analítico actualizador)
    doc.add_heading("ESCUELA PRIMARIA “REVOLUCIÓN SOCIAL DE 1910”", level=1)
    doc.add_paragraph("CCT: 10EPR0361N | Zona: 49 | Sector: 02\nGrado: 6° Grupo “A” | Maestra: Aday Domínguez Domínguez")
    doc.add_paragraph(f"📌 [MODO ESTRICTO ACTIVO • Base seleccionada: {fuente}]")
    doc.add_paragraph("---")
    
    # Datos generales
    doc.add_heading(f"Campo Formativo: {campo}", level=2)
    doc.add_paragraph(f"Proyecto Oficial vinculado (CONALITEG): {proyecto}")
    doc.add_paragraph(f"Temporalidad: {periodo}")
    
    # Desarrollo según la fuente seleccionada sin alteraciones
    if "Solo Programa Analítico" in fuente:
        doc.add_heading("1. Diagnóstico, Problemática Local y Contextualización", level=3)
        doc.add_paragraph("• Basado estrictamente en el Programa Analítico de la escuela (Dificultades lectoras, ortografía, fluidez y asuntos comunitarios de Durango).")
        doc.add_heading("2. Secuencia Didáctica basada en el Programa Analítico", level=3)
        doc.add_paragraph("• Respeto absoluto a la dosificación y contenidos locales registrados para las semanas seleccionadas.")
        
    elif "Solo Programa Sintético" in fuente:
        doc.add_heading("1. Contenidos y PDA Oficiales (Programa Sintético Fase 5)", level=3)
        doc.add_paragraph("• Extracción literal de los Procesos de Desarrollo de Aprendizaje (PDA) de la SEP sin modificaciones de contexto.")
        doc.add_heading("2. Secuencia Didáctica Nacional", level=3)
        doc.add_paragraph("• Abordaje metodológico general de los Libros de Texto Gratuito.")
        
    else: # Ambos
        doc.add_heading("1. Integración Curricular (Programa Sintético + Programa Analítico)", level=3)
        doc.add_paragraph("• Fusión rigurosa: Se cruzan los PDA oficiales nacionales con las problemáticas contextualizadas de la escuela primaria.")
        doc.add_heading("2. Secuencia Didáctica Cruzada", level=3)
        doc.add_paragraph("• Actividades alineadas a los saberes de la comunidad y los ejes articuladores de la Nueva Escuela Mexicana.")

    doc.add_heading("3. Evaluación Formativa", level=3)
    doc.add_paragraph("• Valoración continua mediante rúbricas y seguimiento estricto del proceso de aprendizaje.")
    
    # Guardar en memoria
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

# 3. Botón de descarga
nombre_archivo = f"Planeacion_Estricta_{campo_form.replace(' ', '_')}.docx"

docx_buffer = generar_documento_estricto(campo_form, proyecto_ltg, periodo_val, modo_estricto_fuente)

st.download_button(
    label="📥 Descargar Planeación Estricta en Word (.docx)",
    data=docx_buffer,
    file_name=nombre_archivo,
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    type="primary",
    use_container_width=True
)
